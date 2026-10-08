from __future__ import annotations

import json
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
import pytest

from app.core.config import settings
from app.main import app
from app.retrieval.medlineplus import MedlinePlusRetriever
from app.services.report_upload import ReportExtractionError, extract_report
from app.services.clinical_structure import structure_medical_report
from scripts.build_medlineplus_snapshot import build_snapshot


client = TestClient(app)


def test_txt_upload_returns_verbatim_text_and_does_not_store_file() -> None:
    source = "  Acute myocardial infarction.\nNo pleural effusion.  "
    response = client.post(
        "/api/report/upload",
        files={"file": ("report.txt", source.encode("utf-8"), "text/plain")},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["filename"] == "report.txt"
    assert payload["source_type"] == "TXT"
    assert payload["file_size_bytes"] == len(source.encode("utf-8"))
    assert payload["extracted_text"] == source
    assert payload["page_count"] is None
    assert payload["status"] == "extracted"
    assert payload["report_id"]

    comparison = client.post(
        "/api/compare",
        json={"report": payload["extracted_text"], "top_k": 3},
    )
    assert comparison.status_code == 200
    assert comparison.json()["report"] == source
    assert comparison.json()["b1"]["evidence"] == []
    assert comparison.json()["b2"]["evidence"] == comparison.json()["b3"]["evidence"]


def test_docx_extraction_is_local_and_preserves_paragraphs() -> None:
    document = Document()
    document.add_paragraph("Mild cardiomegaly.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Finding"
    table.cell(0, 1).text = "No pleural effusion."
    document.add_paragraph("No pleural effusion.")
    content = BytesIO()
    document.save(content)

    extracted = extract_report("report.docx", content.getvalue())

    assert extracted.source_type == "DOCX"
    assert extracted.text == "Mild cardiomegaly.\nFinding | No pleural effusion.\nNo pleural effusion."


def test_pdf_text_is_extracted_locally() -> None:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
    })
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 72 720 Td (Mild cardiomegaly.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(content)
    pdf = BytesIO()
    writer.write(pdf)

    extracted = extract_report("report.pdf", pdf.getvalue())

    assert extracted.source_type == "PDF"
    assert extracted.text.strip() == "Mild cardiomegaly."
    assert extracted.page_count == 1

    response = client.post(
        "/api/report/upload",
        files={"file": ("report.pdf", pdf.getvalue(), "application/pdf")},
    )
    assert response.status_code == 200
    assert response.json()["page_count"] == 1


def test_pdf_page_boundaries_reach_structured_result_provenance() -> None:
    writer = PdfWriter()
    for text in ("ApoB = 46.00 mg/dL", "High-sensitivity Troponin-I = 4 ng/L"):
        page = writer.add_blank_page(width=612, height=792)
        font = DictionaryObject({
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        })
        page[NameObject("/Resources")] = DictionaryObject({
            NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)}),
        })
        content = DecodedStreamObject()
        page_text = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        content.set_data(page_text.encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(content)

    pdf = BytesIO()
    writer.write(pdf)
    extracted = extract_report("report.pdf", pdf.getvalue())
    structured = structure_medical_report(extracted.text)
    results = {item.test_name: item for item in structured.results}

    assert extracted.page_count == 2
    assert "\f" in extracted.text
    assert results["Apolipoprotein B"].source_page == 1
    assert results["High-sensitivity Troponin-I"].source_page == 2
    assert results["High-sensitivity Troponin-I"].source_span.page == 2


def test_pdf_upload_rejects_unreadable_font_glyphs_before_analysis(monkeypatch) -> None:
    monkeypatch.setattr(
        "app.services.report_upload._extract_pdf",
        lambda _: "□" * 14_884,
    )

    response = client.post(
        "/api/report/upload",
        files={"file": ("Z615.pdf", b"synthetic-pdf-content", "application/pdf")},
    )

    assert response.status_code == 422
    assert "could not be decoded into readable report text" in response.json()["detail"]
    assert "It was not analyzed" in response.json()["detail"]


def test_compare_rejects_unreadable_report_text() -> None:
    response = client.post(
        "/api/compare",
        json={"report": "□" * 14_884, "top_k": 3},
    )

    assert response.status_code == 422
    assert "does not appear readable" in response.json()["detail"]


def test_extracted_text_readability_accepts_non_latin_report_text() -> None:
    extracted = extract_report("report.txt", "रिपोर्ट में कोई समस्या नहीं है।".encode("utf-8"))

    assert extracted.text == "रिपोर्ट में कोई समस्या नहीं है।"


def test_upload_rejects_empty_and_unsupported_files() -> None:
    empty = client.post("/api/report/upload", files={"file": ("empty.txt", b"", "text/plain")})
    unsupported = client.post("/api/report/upload", files={"file": ("report.csv", b"x", "text/csv")})

    assert empty.status_code == 422
    assert "empty" in empty.json()["detail"].lower()
    assert unsupported.status_code == 415
    assert "Unsupported file type" in unsupported.json()["detail"]


def test_local_ocr_error_is_explicit_when_tesseract_is_missing(monkeypatch) -> None:
    monkeypatch.setattr("app.services.report_upload.shutil.which", lambda _: None)

    with pytest.raises(ReportExtractionError, match="Local OCR is not available"):
        extract_report("scan.png", b"not-a-real-image")


def test_upload_enforces_configured_size_limit(monkeypatch) -> None:
    monkeypatch.setattr(settings, "max_upload_bytes", 4)

    response = client.post(
        "/api/report/upload",
        files={"file": ("report.txt", b"12345", "text/plain")},
    )

    assert response.status_code == 413
    assert "Maximum upload size" in response.json()["detail"]


def test_pdf_extraction_failure_returns_actionable_error() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    pdf = BytesIO()
    writer.write(pdf)

    with pytest.raises(ReportExtractionError, match="No text could be extracted"):
        extract_report("scan.pdf", pdf.getvalue())


def test_findings_preserve_negative_wording() -> None:
    findings = MedlinePlusRetriever.report_findings(
        "Mild cardiomegaly with bilateral pleural effusion. No pleural effusion. No evidence of pneumothorax."
    )

    assert [finding["text"] for finding in findings] == [
        "Mild cardiomegaly",
        "bilateral pleural effusion.",
        "No pleural effusion.",
        "No evidence of pneumothorax.",
    ]
    assert [finding["polarity"] for finding in findings] == [
        "AFFIRMED",
        "AFFIRMED",
        "NEGATED",
        "NEGATED",
    ]


def test_local_snapshot_search_uses_positive_report_text_and_preserves_metadata(tmp_path) -> None:
    snapshot = tmp_path / "medlineplus.jsonl"
    snapshot.write_text(
        "\n".join(
            [
                '{"document_id":"ami","title":"Myocardial infarction","text":"Myocardial infarction and heart attack information.","url":"https://medlineplus.gov/heartattack.html","snapshot_date":"2026-01-01"}',
                '{"document_id":"effusion","title":"Pleural effusion","text":"Pleural effusion information.","url":"https://example.test/effusion","snapshot_date":"2026-01-01"}',
                '{"document_id":"unrelated","title":"Atelectasis","text":"Atelectasis information.","url":"https://example.test/atelectasis","snapshot_date":"2026-01-01"}',
            ]
        ),
        encoding="utf-8",
    )
    retriever = MedlinePlusRetriever()
    retriever.load_snapshot(snapshot)

    evidence = retriever.local_search(
        "Acute myocardial infarction involving the inferior wall. No pleural effusion.",
        top_k=5,
    )

    assert [item.document_id for item in evidence] == ["ami"]
    assert evidence[0].url == "https://medlineplus.gov/heartattack.html"
    assert evidence[0].snapshot_date == "2026-01-01"


def test_official_xml_snapshot_converter_preserves_retrieval_metadata(tmp_path) -> None:
    source = tmp_path / "health-topics.xml"
    source.write_text(
        """<?xml version="1.0" encoding="UTF-8"?>
<health-topics>
  <health-topic id="100" url="https://medlineplus.gov/heartattack.html" language="English">
    <title>Heart Attack</title>
    <also-called>Myocardial Infarction</also-called>
    <full-summary>&lt;p&gt;A heart attack affects blood flow to the heart.&lt;/p&gt;</full-summary>
  </health-topic>
</health-topics>
""",
        encoding="utf-8",
    )
    destination = tmp_path / "snapshot.jsonl"

    count = build_snapshot(source, destination, "2026-01-15")
    record = json.loads(destination.read_text(encoding="utf-8"))

    assert count == 1
    assert record == {
        "document_id": "100",
        "title": "Heart Attack",
        "text": "Heart Attack Myocardial Infarction A heart attack affects blood flow to the heart.",
        "url": "https://medlineplus.gov/heartattack.html",
        "snapshot_date": "2026-01-15",
    }
